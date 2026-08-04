from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from bs4 import BeautifulSoup


def read_document(file_path: str, file_type: str) -> str:
    """Đọc nội dung text thuần từ 1 file, tùy theo định dạng. Output thống
    nhất là text thuần — các bước sau (chunking, embedding) không cần biết
    file gốc là định dạng gì nữa."""
    reader = _READERS.get(file_type)
    if reader is None:
        raise ValueError(f"Không hỗ trợ định dạng: {file_type}")
    return reader(file_path)


def _read_md(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8")


def _read_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages_text)


def _read_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Bảng trong docx KHÔNG nằm trong doc.paragraphs — phải trích riêng,
    # nếu không sẽ mất toàn bộ nội dung bảng (như bảng phân loại rủi ro
    # thiết bị trong tài liệu BYOD mẫu).
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def _read_html(file_path: str) -> str:
    html = Path(file_path).read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()  # bỏ code JS/CSS nếu có, tránh lẫn vào nội dung
    return soup.get_text(separator="\n", strip=True)


_READERS = {
    "md": _read_md,
    "pdf": _read_pdf,
    "docx": _read_docx,
    "html": _read_html,
}